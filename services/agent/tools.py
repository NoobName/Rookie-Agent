import asyncio
import difflib
import hashlib
import json
import os
import time
from pathlib import Path

import psutil
from pydantic import Field

from .document_edit import DocumentEdit, edit_document
from .extensions import BrowserArgs, BrowserTools, DesktopArgs, DesktopTools
from .integrations import EmailArgs, SearchArgs, search, send_email
from .mcp_registry import resolve as resolve_mcp
from .models import AcceptanceCheck, Command, FileRead, FileWrite, GitCommit, Revert, Strict
from .policy import inside
from .store import uid
from .task_plan import TaskPlan
from .workspace import (
    DirectoryList,
    FileEdit,
    WorkspaceInspect,
    directory_listing,
    edit_file,
    inspect_workspace,
)


class Empty(Strict):
    pass


class FileSearch(Strict):
    query: str = Field(min_length=1, max_length=200)
    path: str = "."
    max_results: int = Field(default=50, ge=1, le=200)


class Document(Strict):
    path: str
    content: str | None = None


class Fetch(Strict):
    url: str


class MCPCall(Strict):
    server_id: str = Field(pattern=r"^[a-zA-Z0-9_-]{1,80}$")
    server_sha256: str | None = None
    tool: str = ""
    arguments: dict = {}


SCHEMAS = {
    "files.read": FileRead,
    "files.write": FileWrite,
    "files.list": DirectoryList,
    "workspace.inspect": WorkspaceInspect,
    "files.edit": FileEdit,
    "task.plan": TaskPlan,
    "terminal.run": Command,
    "document.edit": DocumentEdit,
    "files.search": FileSearch,
    "git.diff": Empty,
    "document.process": Document,
    "web.fetch": Fetch,
    "mcp.call": MCPCall,
    "browser.action": BrowserArgs,
    "desktop.action": DesktopArgs,
    "web.search": SearchArgs,
    "email.send": EmailArgs,
    "files.revert": Revert,
    "git.status": Empty,
    "git.commit": GitCommit,
    "checks.run": AcceptanceCheck,
}
DESCRIPTIONS = {
    "files.search": "Literal text search inside workspace; bounded to 1000 files, skips symlinks and dependency/private directories. Returns path, line and excerpt.",
    "document.edit": "Edit an existing document using source expected_sha256 from document.process. Omit output or set it equal to source for in-place editing; a different output may be overwritten. Optional expected_output_sha256 checks its version. DOCX run-preserving replacements/append, XLSX typed cell edits preserving other formulas, PDF 1-based page extraction/reorder, text replacements. Edits are validated before atomic replacement; concurrent changes abort the write.",
    "files.read": "Read a bounded UTF-8 page within workspace. offset and max_bytes are byte offsets. Follow next_offset while truncated. sha256 covers the entire file for overwrite checks. Binary files return metadata, not text; inspect images with a suitable image tool or verify via terminal.",
    "files.write": "Create or replace UTF-8 file. Existing file requires matching expected_sha256.",
    "files.list": "List workspace directory recursively with bounded depth and entry count.",
    "workspace.inspect": "Inspect project tree, build manifests, package scripts and installed tool paths without executing commands. Read relevant source next.",
    "files.edit": "Precisely edit existing UTF-8 source using old/new replacements and expected_sha256 from files.read. All replacements must match expected_matches before any write. Preserves CRLF and untouched content; returns diff and new hash.",
    "task.plan": "Maintain working steps, progress summary and next action across requests and resumes. Does not change user requirements, permissions or acceptance checks.",
    "terminal.run": "Execute argv WITHOUT OS sandbox. Can access outside workspace. Requires explicit authorization.",
    "git.diff": "Read git diff; no commits.",
    "document.process": "Read or create txt/md/json/csv/docx/xlsx/pdf. Workspace read/write follows the selected approval mode.",
    "web.fetch": "Fetch public HTTP(S) page, preserving source URL and retrieval time. No search provider included.",
    "mcp.call": "Use a user-registered MCP server_id to discover or call tools. Launch commands and URLs can only come from saved server settings; never supply your own. Requires authorization.",
    "browser.action": "Independent browser. Observe before each action. Submit requires separate action and explicit approval.",
    "desktop.action": "Windows UIA. Human must arm control. Each input requires fresh observation, stable target and foreground. Ctrl+Alt+F12 emergency stop.",
    "web.search": "Query the user-configured real search API and preserve sources. Fails clearly when unconfigured.",
    "email.send": "Send an external email. Disabled unless user configured SMTP. Always requires exact recipient/body approval.",
    "files.revert": "Revert only a recorded change from this task, and only if the current file still matches the written hash. Requires explicit approval.",
    "git.status": "Read Git working tree status.",
    "git.commit": "Commit ONLY explicitly listed paths; may execute repository hooks, requires independent execution authorization.",
}
DESCRIPTIONS["checks.run"] = (
    "Run an exact user-planned acceptance check; models cannot change the approved acceptance criteria."
)


def schemas(names):
    return [
        {
            "type": "function",
            "function": {
                "name": n.replace(".", "__"),
                "description": DESCRIPTIONS[n],
                "parameters": SCHEMAS[n].model_json_schema(),
            },
        }
        for n in names
        if n in SCHEMAS and n != "checks.run"
    ]


def kill_tree(pid):
    try:
        parent = psutil.Process(pid)
    except psutil.NoSuchProcess:
        return
    children = parent.children(recursive=True)
    for p in reversed(children):
        try:
            p.kill()
        except psutil.NoSuchProcess:
            pass
    try:
        parent.kill()
    except psutil.NoSuchProcess:
        pass
    psutil.wait_procs(children + [parent], timeout=3)


class Tools:
    def __init__(self, store=None):
        self.store = store
        self.processes = {}
        self.browser = BrowserTools()
        self.desktop = DesktopTools()
        if store:
            saved = store.rows("SELECT value FROM settings WHERE key='browser_channel'")
            if saved:
                channel = json.loads(saved[0]["value"])
                self.browser.channel = None if channel == "chromium" else channel

    async def synchronous(self, fn, *args):
        # Keep the single executor occupied until an uninterruptible side effect settles.
        job = asyncio.create_task(asyncio.to_thread(fn, *args))
        try:
            return await asyncio.shield(job)
        except asyncio.CancelledError:
            await asyncio.gather(job, return_exceptions=True)
            raise

    async def execute(self, task, name, args, emit):
        name = name.replace("__", ".")
        validated = SCHEMAS[name].model_validate(args).model_dump()
        root = task["spec"]["workspace"]
        start = time.time()
        result = {"status": "succeeded", "started": start, "exit_code": None, "artifacts": []}
        try:
            if name == "files.list":
                result.update(await self.synchronous(directory_listing, root, validated))
            elif name == "workspace.inspect":
                result.update(await self.synchronous(inspect_workspace, root, validated))
            elif name == "files.edit":
                result.update(
                    await self.synchronous(edit_file, root, validated, self.store, task["id"])
                )
            elif name == "task.plan":
                await emit("task_plan", validated)
                result["output"] = validated
            elif name == "files.search":
                folder = inside(root, validated["path"])
                matches = []
                count = 0
                truncated = False
                for current, dirs, files in os.walk(folder, followlinks=False):
                    dirs[:] = [
                        d
                        for d in dirs
                        if d
                        not in (
                            ".git",
                            "node_modules",
                            ".runtime",
                            ".venv",
                            ".agent-browser-profile",
                            "data",
                        )
                        and not Path(current, d).is_symlink()
                    ]
                    for filename in files:
                        candidate = Path(current, filename)
                        if candidate.is_symlink():
                            continue
                        candidate = inside(root, str(candidate))
                        count += 1
                        if count > 1000:
                            truncated = True
                            break
                        if candidate.stat().st_size > 1_000_000:
                            continue
                        try:
                            lines = candidate.read_text("utf-8").splitlines()
                        except (UnicodeError, OSError):
                            continue
                        for index, line in enumerate(lines, 1):
                            if validated["query"] in line:
                                matches.append(
                                    {
                                        "path": str(candidate.relative_to(Path(root).resolve())),
                                        "line": index,
                                        "text": line[:500],
                                    }
                                )
                            if len(matches) >= validated["max_results"]:
                                truncated = True
                                break
                        if truncated:
                            break
                    if truncated:
                        break
                result.update(
                    output=matches,
                    truncated=truncated,
                    verification="字面匹配；忽略二进制、超过 1 MB 的文件和依赖/私有目录",
                )
            elif name == "files.read":
                import codecs

                p = inside(root, validated["path"])
                digest = hashlib.sha256()
                with p.open("rb") as handle:
                    size = 0
                    while chunk := handle.read(1024 * 1024):
                        digest.update(chunk)
                        size += len(chunk)
                    handle.seek(0)
                    head = handle.read(4096)
                    handle.seek(validated["offset"])
                    raw = handle.read(validated["max_bytes"])
                result.update(
                    sha256=digest.hexdigest(), size_bytes=size, offset=validated["offset"]
                )
                if b"\0" in head or p.suffix.lower() in (
                    ".exe",
                    ".dll",
                    ".bmp",
                    ".png",
                    ".jpg",
                    ".jpeg",
                    ".pdf",
                    ".zip",
                    ".docx",
                    ".xlsx",
                ):
                    result.update(
                        output="二进制文件：已返回大小和哈希。不能作为 UTF-8 源码读取；请使用对应文档/图像工具，或终端验证文件格式。",
                        binary=True,
                        truncated=False,
                    )
                else:
                    decoder = codecs.getincrementaldecoder("utf-8")()
                    try:
                        text = decoder.decode(raw, final=validated["offset"] + len(raw) >= size)
                        next_offset = validated["offset"] + len(raw) - len(decoder.getstate()[0])
                        result.update(
                            output=text,
                            binary=False,
                            truncated=next_offset < size,
                            next_offset=next_offset if next_offset < size else None,
                        )
                    except UnicodeDecodeError:
                        result.update(
                            output="文件不是有效 UTF-8 文本，或 offset 不在字符边界。请用上次返回的 next_offset，或通过终端按实际编码读取。",
                            binary=True,
                            truncated=False,
                        )
            elif name == "files.write":
                p = inside(root, validated["path"])
                old = p.read_bytes() if p.exists() else None
                if (
                    old is not None
                    and hashlib.sha256(old).hexdigest() != validated["expected_sha256"]
                ):
                    raise ValueError("文件已经存在或发生变化；先读取并提交匹配的 expected_sha256")
                if old is not None and len(old) > 1_000_000:
                    raise ValueError("文件过大，拒绝文本覆盖")
                old_text = old.decode("utf-8") if old is not None else None
                content = validated["content"]
                p.parent.mkdir(parents=True, exist_ok=True)
                if p.suffix.lower() in (".bat", ".cmd"):
                    content = (
                        content.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\r\n")
                    )
                # Revalidate after creating parents to detect existing junctions.
                inside(root, validated["path"])
                p.write_bytes(content.encode("utf-8"))
                if self.store:
                    change_id = uid()
                    self.store.run(
                        "INSERT INTO changes VALUES(?,?,?,?,?,0)",
                        (
                            change_id,
                            task["id"],
                            str(p),
                            old_text,
                            hashlib.sha256(p.read_bytes()).hexdigest(),
                        ),
                    )
                    result["change_id"] = change_id
                result.update(
                    output=f"已写入 {p.name}",
                    artifacts=[str(p.relative_to(Path(root).resolve()))],
                    diff="".join(
                        difflib.unified_diff(
                            (old_text or "").splitlines(True),
                            content.splitlines(True),
                            fromfile=str(p),
                            tofile=str(p),
                        )
                    ),
                )
            elif name == "files.revert":
                if not self.store:
                    raise ValueError("缺少变更记录存储")
                change = self.store.one(
                    "SELECT * FROM changes WHERE id=? AND task_id=?",
                    (validated["change_id"], task["id"]),
                )
                p = inside(root, change["path"])
                if (
                    change["reverted"]
                    or not p.exists()
                    or hashlib.sha256(p.read_bytes()).hexdigest() != change["after_sha"]
                ):
                    raise ValueError("文件已有后续修改或已回退，不覆盖用户改动")
                if change["before_text"] is None:
                    p.unlink()
                else:
                    p.write_bytes(change["before_text"].encode("utf-8"))
                self.store.run("UPDATE changes SET reverted=1 WHERE id=?", (change["id"],))
                result["output"] = "仅回退了本任务记录的变更"
            elif name in ("terminal.run", "git.diff", "git.status", "git.commit"):
                if name == "git.commit":
                    paths = [
                        str(inside(root, p).relative_to(Path(root).resolve()))
                        for p in validated["paths"]
                    ]
                    argv = ["git", "commit", "--only", "-m", validated["message"], "--", *paths]
                elif name == "git.status":
                    argv = ["git", "status", "--porcelain=v1"]
                else:
                    argv = validated.get(
                        "argv", ["git", "--no-pager", "diff", "--no-ext-diff", "--no-textconv"]
                    )
                cwd = str(inside(root, validated["cwd"])) if validated.get("cwd") else root
                result.update(
                    await self.command(
                        task["id"],
                        argv,
                        cwd,
                        validated.get("timeout", 30),
                        emit,
                        validated.get("encoding", "utf-8"),
                    )
                )
            elif name == "document.process":
                result.update(await self.synchronous(document, root, validated))
            elif name == "document.edit":
                result.update(await self.synchronous(edit_document, root, validated))
            elif name == "web.fetch":
                import ipaddress
                import socket
                from urllib.parse import urlparse

                import httpx

                parsed = urlparse(validated["url"])
                if parsed.scheme not in ("http", "https") or not parsed.hostname:
                    raise ValueError("仅支持 HTTP(S)")
                addresses = await asyncio.to_thread(socket.getaddrinfo, parsed.hostname, None)
                if any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
                    raise PermissionError("禁止访问内网地址")
                async with httpx.AsyncClient(timeout=30, follow_redirects=False) as client:
                    async with client.stream("GET", validated["url"]) as response:
                        response.raise_for_status()
                        chunks = []
                        size = 0
                        async for chunk in response.aiter_bytes():
                            size += len(chunk)
                            if size > 2_000_000:
                                raise ValueError("网页超出大小限制")
                            chunks.append(chunk)
                result.update(
                    output=b"".join(chunks).decode("utf-8", errors="replace"),
                    source={"url": validated["url"], "retrieved": time.time()},
                )
            elif name == "mcp.call":
                result.update(
                    output=await asyncio.wait_for(mcp_call(resolve_mcp(self.store, validated)), 60)
                )
            elif name == "browser.action":
                result.update(await self.browser.run(task, validated))
            elif name == "desktop.action":
                if validated["action"] == "screenshot":
                    output = inside(root, validated.get("output") or "")
                    if output.exists():
                        raise ValueError("截图不覆盖已有文件")
                    output.parent.mkdir(parents=True, exist_ok=True)
                    inside(root, str(output))
                    validated["output"] = str(output)
                result.update(await self.synchronous(self.desktop.run, validated))
                if validated["action"] == "screenshot":
                    result["artifacts"] = [str(output.relative_to(Path(root).resolve()))]
            elif name == "web.search":
                result.update(await search(self.store, validated))
            elif name == "email.send":
                result.update(await self.synchronous(send_email, self.store, validated))
            elif name == "checks.run":
                from .verification import perform_check

                result.update(await perform_check(self, task, validated, emit))
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            error = (
                f"{type(exc).__name__}: {exc}"
                if name not in ("mcp.call", "email.send", "web.search")
                else f"{type(exc).__name__}：外部服务调用失败，请检查连接、授权和配置；不记录远程异常正文。"
            )
            result.update(status="failed", error=error)
        result["finished"] = time.time()
        return result

    async def command(self, task_id, argv, cwd, timeout, emit, encoding="auto"):
        import subprocess

        process = await asyncio.create_subprocess_exec(
            *argv,
            cwd=cwd,
            stdin=asyncio.subprocess.DEVNULL,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
            **({"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {}),
        )
        self.processes[task_id] = process
        captured = {"stdout": "", "stderr": ""}
        total = 0

        class OutputLimit(Exception):
            pass

        async def pump(pipe, key):
            import codecs

            nonlocal total
            current = "utf-8" if encoding == "auto" else encoding
            decoder = codecs.getincrementaldecoder(current)(
                errors="strict" if encoding == "auto" else "replace"
            )
            while chunk := await pipe.read(4096):
                total += len(chunk)
                if total > 2_000_000:
                    raise OutputLimit()
                pending = decoder.getstate()[0]
                try:
                    text = decoder.decode(chunk)
                except UnicodeDecodeError:
                    if os.name == "nt":
                        import ctypes

                        current = "cp" + str(ctypes.windll.kernel32.GetACP())
                    else:
                        current = "utf-8"
                    decoder = codecs.getincrementaldecoder(current)(errors="replace")
                    text = decoder.decode(pending + chunk)
                captured[key] = (captured[key] + text)[-30000:]
                await emit(key, {"text": text})
            try:
                tail = decoder.decode(b"", final=True)
            except UnicodeDecodeError:
                tail = decoder.getstate()[0].decode(current, errors="replace")
            if tail:
                captured[key] = (captured[key] + tail)[-200000:]
                await emit(key, {"text": tail})

        try:
            await asyncio.wait_for(
                asyncio.gather(
                    pump(process.stdout, "stdout"), pump(process.stderr, "stderr"), process.wait()
                ),
                timeout,
            )
            return {
                **captured,
                "exit_code": process.returncode,
                "status": "succeeded" if process.returncode == 0 else "failed",
            }
        except (asyncio.TimeoutError, OutputLimit) as exc:
            await asyncio.to_thread(kill_tree, process.pid)
            await process.wait()
            return {
                **captured,
                "exit_code": process.returncode,
                "status": "failed",
                "timeout_seconds": timeout,
                "error": (
                    f"执行超时（{timeout:g} 秒），已终止本次进程树；检查交互等待、脚本循环和输出错误。"
                    if isinstance(exc, asyncio.TimeoutError)
                    else "输出超过 2 MB，已终止进程树；请检查脚本循环或重复报错，勿直接重试。"
                ),
            }
        except asyncio.CancelledError:
            await asyncio.to_thread(kill_tree, process.pid)
            await process.wait()
            raise
        finally:
            self.processes.pop(task_id, None)


def document(root, args):
    p = inside(root, args["path"])
    content = args["content"]
    ext = p.suffix.lower()
    if content is not None:
        if p.exists():
            raise ValueError("文档创建不覆盖已有文件，请使用新的输出路径")
        p.parent.mkdir(parents=True, exist_ok=True)
    if ext == ".docx":
        from docx import Document

        if content is not None:
            d = Document()
            d.add_paragraph(content)
            d.save(p)
        output = "\n".join(x.text for x in Document(p).paragraphs)
    elif ext == ".xlsx":
        from openpyxl import Workbook, load_workbook

        if content is not None:
            w = Workbook()
            for row in json.loads(content):
                w.active.append(row)
            w.save(p)
        w = load_workbook(p, data_only=False)
        output = list(w.active.values)
        w.close()
    elif ext == ".pdf":
        from pypdf import PdfReader

        if content is not None:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen.canvas import Canvas

            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            c = Canvas(str(p))
            c.setFont("STSong-Light", 12)
            y = 800
            for line in content.splitlines():
                for at in range(0, max(1, len(line)), 38):
                    c.drawString(40, y, line[at : at + 38])
                    y -= 18
                    if y < 50:
                        c.showPage()
                        c.setFont("STSong-Light", 12)
                        y = 800
            c.save()
        output = "\n".join(page.extract_text() or "" for page in PdfReader(p).pages)
        if not output.strip():
            output = "未提取到文字：可能为扫描 PDF，未配置 OCR。"
    elif ext in (".txt", ".md", ".json", ".csv"):
        if content is not None:
            p.write_text(content, encoding="utf-8")
        output = p.read_text(encoding="utf-8")
        if ext == ".json":
            json.loads(output)
    else:
        raise ValueError("不支持的文档格式")
    return {
        "output": output,
        "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
        "artifacts": [str(p.relative_to(Path(root).resolve()))] if content is not None else [],
        "verification": "文件重新解析成功；未进行布局渲染、公式重算或 Office 保真验证",
    }


async def mcp_call(args):
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    from mcp.client.streamable_http import streamablehttp_client

    if args["transport"] == "stdio":
        transport = stdio_client(StdioServerParameters(command=args["command"], args=args["args"]))
    elif args["transport"] == "http":
        transport = streamablehttp_client(args["url"], headers=args.get("_headers"))
    else:
        raise ValueError("MCP transport 必须是 stdio 或 http")
    async with transport as streams:
        async with ClientSession(streams[0], streams[1]) as session:
            await session.initialize()
            discovered = await session.list_tools()
            if not args["tool"]:
                return discovered.model_dump(mode="json")
            names = {t.name for t in discovered.tools}
            if args["tool"] not in names:
                raise ValueError("MCP 工具不存在")
            result = await session.call_tool(args["tool"], args["arguments"])
            output = result.model_dump(mode="json")
    if result.isError:
        raise RuntimeError("MCP 服务返回工具错误")
    return output
