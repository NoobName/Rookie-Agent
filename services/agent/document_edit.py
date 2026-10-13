"""Version-checked document edits committed with atomic replacement."""

import hashlib
import json
import os
import re
import tempfile
from io import BytesIO
from pathlib import Path

from pydantic import Field

from .models import Strict
from .policy import inside


class Replacement(Strict):
    old: str = Field(min_length=1, max_length=10000)
    new: str = Field(max_length=10000)


class CellEdit(Strict):
    sheet: str
    cell: str = Field(pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
    value: str | int | float | bool | None


class DocumentEdit(Strict):
    source: str
    output: str | None = None
    expected_output_sha256: str | None = Field(default=None, pattern=r"^[a-f0-9]{64}$")
    expected_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    replacements: list[Replacement] = Field(default_factory=list, max_length=100)
    append_paragraphs: list[str] = Field(default_factory=list, max_length=100)
    cells: list[CellEdit] = Field(default_factory=list, max_length=500)
    pages: list[int] = Field(default_factory=list, max_length=500)


def _render_document(root, args):
    source = inside(root, args["source"])
    output = inside(root, args["output"])
    if source.suffix.lower() != output.suffix.lower():
        raise ValueError("输入和输出格式必须一致")
    if source.stat().st_size > 20_000_000:
        raise ValueError("文档超过 20 MB 修改限制")
    raw = source.read_bytes()
    if hashlib.sha256(raw).hexdigest() != args["expected_sha256"]:
        raise ValueError("输入文件已变化，请重新读取并核对版本")
    ext = source.suffix.lower()
    changed = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    inside(root, args["output"])
    # All validation and in-memory edits finish before writing output.
    if ext == ".docx":
        if args["cells"] or args["pages"]:
            raise ValueError("DOCX 不支持单元格坐标/PDF 页参数")
        from docx import Document

        doc = Document(BytesIO(raw))

        def paragraphs(container):
            yield from container.paragraphs
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from paragraphs(cell)

        blocks = list(paragraphs(doc))
        for section in doc.sections:
            blocks.extend(paragraphs(section.header))
            blocks.extend(paragraphs(section.footer))
        for replace in args["replacements"]:
            count = 0
            for paragraph in blocks:
                for run in paragraph.runs:
                    n = run.text.count(replace["old"])
                    if n:
                        run.text = run.text.replace(replace["old"], replace["new"])
                        count += n
            if not count:
                raise ValueError("未找到可保持格式的完整文本片段；跨格式片段替换不支持")
            changed += count
        for text in args["append_paragraphs"]:
            doc.add_paragraph(text)
            changed += 1
        if not changed:
            raise ValueError("没有修改内容")
        doc.save(output)
        Document(output)
    elif ext == ".xlsx":
        if args["replacements"] or args["append_paragraphs"] or args["pages"]:
            raise ValueError("XLSX 只接受 cells 修改")
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(raw), data_only=False)
        try:
            if not args["cells"]:
                raise ValueError("缺少单元格修改")
            for edit in args["cells"]:
                if edit["sheet"] not in book.sheetnames:
                    raise ValueError("工作表不存在")
                book[edit["sheet"]][edit["cell"]] = edit["value"]
                changed += 1
            book.save(output)
        finally:
            book.close()
        check = load_workbook(output, data_only=False)
        check.close()
    elif ext == ".pdf":
        if args["replacements"] or args["append_paragraphs"] or args["cells"]:
            raise ValueError("PDF 修改仅支持按页提取/重排；不伪装支持文本保真编辑")
        from pypdf import PdfReader, PdfWriter

        reader = PdfReader(BytesIO(raw))
        writer = PdfWriter()
        if not args["pages"] or any(n < 1 or n > len(reader.pages) for n in args["pages"]):
            raise ValueError("页码必须在文档范围内，从 1 开始")
        for n in args["pages"]:
            writer.add_page(reader.pages[n - 1])
            changed += 1
        with output.open("wb") as stream:
            writer.write(stream)
        PdfReader(output)
    elif ext in (".txt", ".md", ".json", ".csv"):
        if args["cells"] or args["pages"]:
            raise ValueError("文本文件仅支持替换与追加")
        text = raw.decode("utf-8-sig")
        for replace in args["replacements"]:
            count = text.count(replace["old"])
            if not count:
                raise ValueError("未找到待替换文本")
            text = text.replace(replace["old"], replace["new"])
            changed += count
        if args["append_paragraphs"]:
            text += "\n" + "\n".join(args["append_paragraphs"])
            changed += len(args["append_paragraphs"])
        if ext == ".json":
            json.loads(text)
        if not changed:
            raise ValueError("没有修改内容")
        output.write_text(text, encoding="utf-8")
    else:
        raise ValueError("不支持的文档修改格式")
    return {
        "output": f"已在新文件完成 {changed} 处操作；输入文件未修改",
        "artifacts": [str(output.relative_to(Path(root).resolve()))],
        "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "verification": "输出已重新解析；未进行公式重算、OCR 或 Office 复杂格式保真验证",
    }


def edit_document(root, args):
    args = DocumentEdit.model_validate(args).model_dump()
    requested = args["output"] or args["source"]
    source = inside(root, args["source"])
    output = inside(root, requested)
    if source.suffix.lower() != output.suffix.lower():
        raise ValueError("输入和输出格式必须一致")
    if source.stat().st_size > 20_000_000:
        raise ValueError("文档超过 20 MB 修改限制")
    digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None
    source_hash = digest(source)
    if source_hash != args["expected_sha256"]:
        raise ValueError("输入文件已变化，请重新读取并核对版本")
    output_hash = digest(output)
    expected = args.get("expected_output_sha256")
    if expected and expected != output_hash:
        raise ValueError("输出文件已变化，请重新读取并核对版本")
    output.parent.mkdir(parents=True, exist_ok=True)
    if inside(root, requested) != output:
        raise PermissionError("输出路径在执行期间发生变化")
    fd, name = tempfile.mkstemp(prefix=".agent-edit-", suffix=output.suffix, dir=output.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        result = _render_document(root, {**args, "output": str(temporary)})
        if inside(root, args["source"]) != source or inside(root, requested) != output:
            raise PermissionError("文档路径在执行期间发生变化")
        if digest(source) != source_hash:
            raise ValueError("输入文件在修改期间发生变化，请重新读取")
        if digest(output) != output_hash:
            raise ValueError("输出文件在修改期间发生变化，请重新读取")
        os.replace(temporary, output)
        result.update(
            output=f"已{'原地修改' if source==output else '更新输出文档'}；输出已重新解析",
            artifacts=[str(output.relative_to(Path(root).resolve()))],
            overwritten=output_hash is not None,
            in_place=source == output,
        )
        return result
    finally:
        temporary.unlink(missing_ok=True)
