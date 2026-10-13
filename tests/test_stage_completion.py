import hashlib
import json
import os
import sys
import time
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient

from services.agent.accounting import MeteredProvider, UsageLimits, totals
from services.agent.app import create_app
from services.agent.document_edit import DocumentEdit, edit_document
from services.agent.models import ProviderConfig
from services.agent.provider_errors import BudgetExceeded
from services.agent.providers import Reply
from services.agent.store import Store, dump
from services.agent.tools import document
from tests.test_execution import approve_until_done


def test_office_modifications_preserve_input_and_formula(tmp_path):
    from docx import Document
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "source.docx"
    doc = Document()
    run = doc.add_paragraph().add_run("old text")
    run.bold = True
    doc.save(source)
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    args = DocumentEdit(
        source="source.docx",
        output="edited.docx",
        expected_sha256=sha,
        replacements=[{"old": "old", "new": "new"}],
    ).model_dump()
    edit_document(tmp_path, args)
    assert Document(tmp_path / "edited.docx").paragraphs[0].runs[0].bold
    assert Document(tmp_path / "edited.docx").paragraphs[0].text == "new text"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == sha
    assert edit_document(tmp_path, args)["overwritten"] is True
    with pytest.raises(ValueError, match="已变化"):
        edit_document(tmp_path, {**args, "output": "bad.docx", "expected_sha256": "0" * 64})
    book = Workbook()
    book.active["A1"] = 2
    book.active["B1"] = "=A1*2"
    book.save(tmp_path / "source.xlsx")
    sha = hashlib.sha256((tmp_path / "source.xlsx").read_bytes()).hexdigest()
    edit_document(
        tmp_path,
        DocumentEdit(
            source="source.xlsx",
            output="edited.xlsx",
            expected_sha256=sha,
            cells=[{"sheet": "Sheet", "cell": "A1", "value": 5}],
        ).model_dump(),
    )
    book = load_workbook(tmp_path / "edited.xlsx")
    assert book.active["A1"].value == 5 and book.active["B1"].value == "=A1*2"
    book.close()
    document(tmp_path, {"path": "source.pdf", "content": "PDF page content"})
    sha = hashlib.sha256((tmp_path / "source.pdf").read_bytes()).hexdigest()
    edit_document(
        tmp_path,
        DocumentEdit(
            source="source.pdf", output="edited.pdf", expected_sha256=sha, pages=[1, 1]
        ).model_dump(),
    )
    from pypdf import PdfReader

    assert len(PdfReader(tmp_path / "edited.pdf").pages) == 2


async def test_metering_cache_reminder_and_hard_limits(tmp_path):
    store = Store(tmp_path / "db")
    requests = []

    def handle(r):
        requests.append(r)
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"role": "assistant", "content": "OK"}}],
                "usage": {
                    "prompt_tokens": 100,
                    "completion_tokens": 10,
                    "prompt_tokens_details": {"cached_tokens": 60},
                },
            },
        )

    cfg = ProviderConfig(
        id="real",
        name="real",
        model="mock",
        streaming=False,
        input_price=10,
        output_price=20,
        cached_input_price=1,
    )
    p = MeteredProvider(cfg, store, httpx.MockTransport(handle))
    store.run(
        "INSERT INTO settings VALUES(?,?)",
        ("usage_limits", UsageLimits(token_reminder=100, hard_tokens=110).model_dump_json()),
    )

    async def emit(_):
        pass

    reply = await p.reply([], [], emit)
    assert reply.usage_id
    assert totals(store)["known_cost"] == pytest.approx(0.00066)
    assert store.one("SELECT * FROM usage_meta")["cached_input_tokens"] == 60
    assert len(store.rows("SELECT * FROM events WHERE type='notification'")) == 1
    with pytest.raises(BudgetExceeded):
        await p.reply([], [], emit)
    assert len(requests) == 1
    store.run(
        "UPDATE settings SET value=? WHERE key=?",
        (UsageLimits(hard_cost=0.0005).model_dump_json(), "usage_limits"),
    )
    with pytest.raises(BudgetExceeded):
        await p.reply([], [], emit)
    assert len(requests) == 1


def test_evaluation_runs_actual_checks_and_keeps_acceptance_separate(tmp_path):
    app = create_app(tmp_path / "db", "t")

    class Provider:
        async def reply(self, messages, tools, emit):
            if any(m["role"] == "tool" for m in messages):
                return Reply("Done")
            return Reply(
                calls=[
                    {
                        "id": "write",
                        "name": "files.write",
                        "arguments": {
                            "path": "solution.py",
                            "content": "def normalize(values):\n    return sorted(set(values))\n",
                        },
                    }
                ]
            )

    app.state.engine.provider = lambda _: Provider()
    with TestClient(app, headers={"Authorization": "Bearer t"}) as c:
        c.put(
            "/api/providers/test",
            json={"id": "test", "name": "test", "model": "test", "tools": True},
        )
        task = c.post("/api/evaluations", json={"provider": "test", "suite": "coding"}).json()
        assert task["state"] == "planned"
        c.post(f"/api/tasks/{task['id']}/start", json={})
        done = approve_until_done(c, task["id"])
        assert done["state"] == "awaiting_feedback"
        result = c.get("/api/evaluations").json()[0]
        assert result["machine_passed"] and not result["user_accepted"]
        assert c.put("/api/usage/limits", json={"hard_tokens": 0}).status_code == 422
        assert c.get("/api/usage").json()["total"]["requests"] == 2
        assert c.get("/api/notifications?latest=true").json()["cursor"] > 0
        c.delete("/api/sessions/" + task["session_id"])
        assert c.get("/api/evaluations").json() == []


async def test_browser_upload_and_handoff_preserves_cookie(tmp_path):
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    from services.agent.extensions import BrowserArgs, BrowserTools

    class Site(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(b'<input type="file" id="upload"><p>handoff fixture</p>')

        def log_message(self, *_):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Site)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    browser = BrowserTools()
    browser.channel = os.getenv("AGENT_BROWSER_CHANNEL") or (
        "msedge" if sys.platform == "win32" else None
    )
    task = {"id": "fixture", "spec": {"workspace": str(tmp_path)}}
    file = tmp_path / "proof.txt"
    file.write_text("fixture")

    async def action(**args):
        return await browser.run(task, BrowserArgs(**args).model_dump())

    try:
        url = f"http://127.0.0.1:{server.server_port}"
        state = await action(action="open", url=url)
        assert state["controls"][0]["id"] == "upload"
        await action(
            action="upload", input="proof.txt", selector="#upload", observation=state["observation"]
        )
        assert (
            await browser.sessions["fixture"]["page"]
            .locator("#upload")
            .evaluate("el=>el.files[0].name")
            == "proof.txt"
        )
        await browser.sessions["fixture"]["context"].add_cookies(
            [{"name": "fixture", "value": "kept", "url": url, "expires": time.time() + 3600}]
        )
        # Real Chromium headed handoff; no external accounts or user browser profiles.
        assert (await browser.handoff(task, True))["human"]
        with pytest.raises(PermissionError):
            await action(action="observe")
        await browser.handoff(task, False)
        assert any(
            c["name"] == "fixture" for c in await browser.sessions["fixture"]["context"].cookies()
        )
        with pytest.raises(ValueError, match="过期"):
            await action(
                action="upload",
                input="proof.txt",
                selector="#upload",
                observation=state["observation"],
            )
    finally:
        await browser.close()
        server.shutdown()
        server.server_close()


async def test_real_remote_mcp_bearer_and_registry_revision(tmp_path, monkeypatch):
    import asyncio
    import socket
    import threading

    import uvicorn
    from mcp.server.fastmcp import FastMCP

    from services.agent.mcp_registry import resolve, snapshot
    from services.agent.tools import mcp_call

    mcp = FastMCP("http-fixture", stateless_http=True)

    @mcp.tool()
    def multiply(a: int, b: int) -> int:
        return a * b

    inner = mcp.streamable_http_app()
    seen = []

    async def remote(scope, receive, send):
        if scope["type"] == "http":
            headers = dict(scope["headers"])
            seen.append(headers.get(b"authorization"))
            if headers.get(b"authorization") != b"Bearer fixture-secret":
                await send({"type": "http.response.start", "status": 401, "headers": []})
                await send({"type": "http.response.body", "body": b"Unauthorized"})
                return
        await inner(scope, receive, send)

    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    sock.listen()
    port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(remote, log_level="error"))
    thread = threading.Thread(target=lambda: server.run(sockets=[sock]), daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        await asyncio.sleep(0.02)
    store = Store(tmp_path / "db")
    config = {
        "id": "remote",
        "name": "remote",
        "enabled": True,
        "transport": "http",
        "command": "",
        "args": [],
        "url": f"http://127.0.0.1:{port}/mcp",
    }
    store.run("INSERT INTO mcp_servers VALUES(?,?,?)", ("remote", dump(config), "fixture-ref"))
    monkeypatch.setattr(
        "services.agent.mcp_registry.keyring.get_password", lambda *_: "fixture-secret"
    )
    args = snapshot(
        store, {"server_id": "remote", "tool": "multiply", "arguments": {"a": 6, "b": 7}}
    )
    try:
        result = await mcp_call(resolve(store, args))
        assert result["content"][0]["text"] == "42"
        assert seen and all(x == b"Bearer fixture-secret" for x in seen)
        store.run(
            "UPDATE mcp_servers SET config=? WHERE id=?",
            (dump({**config, "name": "changed"}), "remote"),
        )
        with pytest.raises(PermissionError, match="配置已变化"):
            resolve(store, args)
        assert "fixture-secret" not in dump(store.rows("SELECT * FROM mcp_servers"))
    finally:
        server.should_exit = True
        await asyncio.to_thread(thread.join, 5)
        sock.close()


async def test_literal_search_and_command_encoding(tmp_path):
    from services.agent.tools import Tools

    root = tmp_path / "workspace"
    root.mkdir()
    (root / "中文.txt").write_text("alpha\n检索目标\n", "utf-8")
    (root / "sub").mkdir()
    tools = Tools()
    task = {"id": "search", "spec": {"workspace": str(root)}}

    async def emit(*_):
        pass

    result = await tools.execute(task, "files.search", {"query": "检索目标"}, emit)
    assert result["output"][0]["line"] == 2
    result = await tools.execute(
        task,
        "terminal.run",
        {
            "argv": [
                sys.executable,
                "-c",
                "import sys;sys.stdout.buffer.write('中文'.encode('gb18030'))",
            ],
            "cwd": "sub",
            "encoding": "gb18030",
        },
        emit,
    )
    assert result["stdout"] == "中文" and result["exit_code"] == 0
    result = await tools.execute(
        task, "terminal.run", {"argv": [sys.executable, "-c", "print(1)"], "cwd": "../"}, emit
    )
    assert result["status"] == "failed"


async def test_reminder_does_not_create_execution_task(tmp_path):
    from services.agent.scheduler import Schedule, Scheduler

    store = Store(tmp_path / "db")
    now = time.time()

    class Queue:
        async def enqueue(self, _):
            raise AssertionError("Reminder must not execute tasks")

    spec = Schedule(
        name="synthetic reminder", kind="reminder", next_run=now - 1, authorized_until=now + 60
    )
    store.run(
        "INSERT INTO schedules VALUES(?,?,?,?)", ("reminder", spec.model_dump_json(), now - 1, 1)
    )
    scheduler = Scheduler(store, Queue())
    await scheduler.tick(now)
    await scheduler.tick(now + 1)
    assert len(store.rows("SELECT * FROM events WHERE type='notification'")) == 1
    assert store.rows("SELECT * FROM tasks") == []


def test_discovery_key_is_independent_of_api_authentication(tmp_path):
    import hmac

    with TestClient(create_app(tmp_path / "db", "fixture-token")) as c:
        health = c.get("/health").json()
        assert set(health) == {"ok", "app", "version"}
        nonce = "a" * 32
        assert c.get("/health", params={"nonce": nonce}).json() == health
        result = c.get("/discovery", params={"nonce": nonce}).json()
        key = (tmp_path / "db" / "discovery.key").read_text()
        expected = hmac.new(
            key.encode(), (nonce + "\n" + result["data_dir"]).encode(), hashlib.sha256
        ).hexdigest()
        assert result["proof"] == expected and "fixture-token" not in json.dumps(result)
        assert (
            result["proof"]
            != hmac.new(
                b"fixture-token", (nonce + "\n" + result["data_dir"]).encode(), hashlib.sha256
            ).hexdigest()
        )
        assert c.get("/discovery?nonce=invalid").status_code == 400
        assert (
            c.get("/api/bootstrap", headers={"Authorization": "Bearer " + key}).status_code == 401
        )
        for route in ("/health", "/discovery?nonce=" + nonce):
            assert c.get(route, headers={"Origin": "https://example.com"}).status_code == 403
