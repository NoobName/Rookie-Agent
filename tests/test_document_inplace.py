import hashlib

import pytest

from services.agent import document_edit as module


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.mark.parametrize("extension", [".docx", ".xlsx", ".pdf", ".txt"])
def test_real_inplace_document_edits(tmp_path, extension):
    path = tmp_path / ("document" + extension)
    if extension == ".docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph("before")
        doc.save(path)
        changes = {"replacements": [{"old": "before", "new": "after"}]}
    elif extension == ".xlsx":
        from openpyxl import Workbook, load_workbook

        book = Workbook()
        book.active["A1"] = 1
        book.active["B1"] = "=A1*2"
        book.save(path)
        book.close()
        changes = {"cells": [{"sheet": "Sheet", "cell": "A1", "value": 7}]}
    elif extension == ".pdf":
        from pypdf import PdfReader, PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        writer.add_blank_page(200, 200)
        with path.open("wb") as f:
            writer.write(f)
        changes = {"pages": [2, 1]}
    else:
        path.write_text("before", encoding="utf-8")
        changes = {"replacements": [{"old": "before", "new": "after"}]}
    old = sha(path)
    result = module.edit_document(
        tmp_path, {"source": path.name, "expected_sha256": old, **changes}
    )
    assert result["in_place"] and result["overwritten"]
    assert result["artifacts"] == [path.name] and result["sha256"] == sha(path)
    if extension == ".docx":
        assert Document(path).paragraphs[0].text == "after"
    elif extension == ".xlsx":
        book = load_workbook(path)
        assert book.active["A1"].value == 7 and book.active["B1"].value == "=A1*2"
        book.close()
    elif extension == ".pdf":
        assert float(PdfReader(path).pages[0].mediabox.width) == 200
    else:
        assert path.read_text() == "after"
    with pytest.raises(ValueError, match="输入文件已变化"):
        module.edit_document(
            tmp_path, {"source": path.name, "output": path.name, "expected_sha256": old, **changes}
        )
    assert not list(tmp_path.glob(".agent-edit-*"))


def test_invalid_edits_and_concurrent_changes_preserve_existing_output(tmp_path, monkeypatch):
    source = tmp_path / "source.json"
    source.write_text('{"n":1}')
    output = tmp_path / "output.json"
    output.write_text('{"old":true}')
    args = {
        "source": source.name,
        "output": output.name,
        "expected_sha256": sha(source),
        "replacements": [{"old": "1", "new": "2"}],
    }
    with pytest.raises(ValueError, match="输出文件已变化"):
        module.edit_document(tmp_path, {**args, "expected_output_sha256": "0" * 64})
    with pytest.raises(ValueError):
        module.edit_document(
            tmp_path, {**args, "replacements": [{"old": "1", "new": "invalid JSON"}]}
        )
    assert output.read_text() == '{"old":true}'
    result = module.edit_document(tmp_path, {**args, "expected_output_sha256": sha(output)})
    assert result["overwritten"] and not result["in_place"] and output.read_text() == '{"n":2}'
    render = module._render_document

    def concurrent(root, arguments):
        result = render(root, arguments)
        output.write_text('{"external":true}')
        return result

    monkeypatch.setattr(module, "_render_document", concurrent)
    with pytest.raises(ValueError, match="输出文件在修改期间发生变化"):
        module.edit_document(tmp_path, args)
    assert output.read_text() == '{"external":true}' and source.read_text() == '{"n":1}'
    assert not list(tmp_path.glob(".agent-edit-*"))
