import hashlib

import pytest
from docx import Document
from fastapi.testclient import TestClient

from services.agent.app import create_app
from services.agent.policy import PolicyEngine
from tests.test_core import wait_state
from tests.test_execution import create


def test_workspace_document_edit_and_overwrite_are_automatic(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("before")
    doc.save(source)
    args = {
        "source": "source.docx",
        "output": "edited.docx",
        "expected_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "append_paragraphs": ["after"],
    }
    app = create_app(tmp_path / "data", "t")
    with TestClient(app, headers={"Authorization": "Bearer t"}) as c:
        # Both creation and replacement use workspace authorization.
        task = create(
            c, tools=["document.edit"], script=[{"name": "document.edit", "arguments": args}] * 2
        )
        spec = {**task["spec"], "workspace": str(tmp_path)}
        assert c.put("/api/tasks/" + task["id"] + "/spec", json=spec).status_code == 200
        c.post("/api/tasks/" + task["id"] + "/start", json={})
        done = wait_state(c, task["id"], {"awaiting_feedback", "awaiting_approval", "failed"})
        assert done["state"] == "awaiting_feedback"
        assert Document(tmp_path / "edited.docx").paragraphs[-1].text == "after"
        calls = c.get("/api/tasks/" + task["id"] + "/calls").json()
        assert [r["status"] for r in calls] == ["succeeded", "succeeded"]
        assert calls[-1]["result"]["overwritten"] is True
        assert not app.state.engine.store.rows(
            "SELECT * FROM events WHERE task_id=? AND type='approval'", (task["id"],)
        )
        policy = PolicyEngine(app.state.engine.store)
        task = app.state.engine.store.task(task["id"])
        task["spec"]["mode"] = "ask"
        assert not policy.allowed(task, "document.edit", args, "test")
        task["spec"]["mode"] = "workspace"
        for field in ("source", "output"):
            with pytest.raises(PermissionError):
                policy.allowed(task, "document.edit", {**args, field: "../outside.docx"}, "test")
